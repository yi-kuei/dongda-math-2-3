# -*- coding: utf-8 -*-
"""
generate_landscape_worksheets.py
東大版 技術高中數學（第三冊）2-3 對數及其運算的意義
JIS B4 橫向雙欄 Word 學習單產生器 (學生自主學習單 + 教師教學評量解答版)
遵循《PLVS-MATH-PACKAGE-SOP v2.0》規範：
- Word 原生 OMML 數學方程式 (<m:oMath>)
- JIS B4 橫向 (364mm x 257mm) 雙欄排版 (含中央分隔線)
- 學生版留白計算區，教師版醒目紅色 (#DC2626) 原生方程式解答
- 自動將產出之 docx 編碼為 base64 寫入 docx_b64.py
"""

import os
import re
import base64
import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from lxml import etree
import latex2mathml.converter

# 尋找微軟 Office 內建之 MML2OMML.XSL
XSL_PATHS = [
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\Office16\MML2OMML.XSL",
]
XSL_FILE = next((p for p in XSL_PATHS if os.path.exists(p)), None)
if XSL_FILE:
    XSLT_DOC = etree.parse(XSL_FILE)
    TRANSFORM = etree.XSLT(XSLT_DOC)
    print(f"[OK] 成功載入 MML2OMML.XSL: {XSL_FILE}")
else:
    TRANSFORM = None
    print("[WARN] 未找到 MML2OMML.XSL")

def latex_to_omml(latex_str, is_teacher=False):
    """將 LaTeX 數學式轉換為 Word 原生 OMML (<m:oMath>) XML 節點"""
    if not TRANSFORM:
        return None
    try:
        clean_latex = latex_str.replace(r'\dfrac', r'\frac')
        clean_latex = clean_latex.replace(r'\,', r' ')
        clean_latex = clean_latex.replace(r'\parallel', r'||')
        clean_latex = clean_latex.replace(r'\approx', r'~')
        
        mathml_str = latex2mathml.converter.convert(clean_latex)
        mml_tree = etree.fromstring(mathml_str.encode('utf-8'))
        omml_tree = TRANSFORM(mml_tree)
        omml_xml = etree.tostring(omml_tree, encoding='utf-8').decode('utf-8')
        omml_xml = re.sub(r'<\?xml[^>]+\?>', '', omml_xml).strip()
        
        if is_teacher:
            omml_xml = re.sub(
                r'(<m:r>|<m:r\s+[^>]*>)',
                r'\1<m:rPr><w:rPr><w:color w:val="DC2626"/></w:rPr></m:rPr>',
                omml_xml
            )
        return parse_xml(omml_xml)
    except Exception as e:
        return None

def add_math_paragraph(paragraph, full_text, is_teacher=False, default_color=None, bold=False):
    """混合純文字與 $...$ 數學式，自動切分並嵌入原生 OMML 節點"""
    parts = re.split(r'(\$[^\$]+\$)', full_text)
    for part in parts:
        if not part:
            continue
        if part.startswith('$') and part.endswith('$') and len(part) > 2:
            latex_content = part[1:-1].strip()
            omml_node = latex_to_omml(latex_content, is_teacher=is_teacher)
            if omml_node is not None:
                paragraph._p.append(omml_node)
            else:
                run = paragraph.add_run(latex_content)
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
                run.bold = bold
                if is_teacher:
                    run.font.color.rgb = RGBColor(220, 38, 38)
                elif default_color:
                    run.font.color.rgb = default_color
        else:
            run = paragraph.add_run(part)
            run.font.name = "微軟正黑體"
            run.font.size = Pt(11)
            run.bold = bold
            if is_teacher:
                run.font.color.rgb = RGBColor(220, 38, 38)
            elif default_color:
                run.font.color.rgb = default_color

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """設定儲存格邊距 (dxa)"""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_box(doc, text, color_rgb=(30, 58, 138), bg_hex="EFF6FF"):
    """加入精緻標題區塊"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"【{text}】")
    run.font.name = "微軟正黑體"
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = RGBColor(*color_rgb)

def add_box(doc, title, content_lines, is_teacher=False, answer_lines=None):
    """加入題目與計算方框"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="2563EB"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    add_math_paragraph(p, title, default_color=RGBColor(30, 58, 138), bold=True)
    
    for line in content_lines:
        lp = cell.add_paragraph()
        lp.paragraph_format.space_after = Pt(2)
        add_math_paragraph(lp, line)
        
    if is_teacher and answer_lines:
        ap = cell.add_paragraph()
        ap.paragraph_format.space_before = Pt(4)
        ap.paragraph_format.space_after = Pt(2)
        arun = ap.add_run("【詳解與結論】：")
        arun.font.name = "微軟正黑體"
        arun.font.size = Pt(10.5)
        arun.bold = True
        arun.font.color.rgb = RGBColor(220, 38, 38)
        for aline in answer_lines:
            alp = cell.add_paragraph()
            alp.paragraph_format.space_after = Pt(2)
            add_math_paragraph(alp, aline, is_teacher=True)
    else:
        for _ in range(4):
            sp = cell.add_paragraph()
            sp.paragraph_format.space_after = Pt(0)
            srun = sp.add_run("　")
            srun.font.size = Pt(11)

def create_worksheet(is_teacher=False):
    doc = docx.Document()
    
    # 1. 版面設定：JIS B4 橫向 (364mm x 257mm)
    section = doc.sections[0]
    section.page_width = Mm(364)
    section.page_height = Mm(257)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)
    
    # 2. 雙欄配置 (含分隔線)
    sectPr = section._sectPr
    cols = OxmlElement('w:cols')
    cols.set(qn('w:num'), '2')
    cols.set(qn('w:sep'), '1')       # 垂直分隔線
    cols.set(qn('w:space'), '720')   # 欄距 720 dxa
    sectPr.append(cols)
    
    # 3. 頁尾：PLVS 數學科  教育學習用
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    status_str = "[教師評量解答版 ‧ 密碼 88888888]" if is_teacher else "[學生自主學習單]"
    f_run = f_p.add_run(f"東大版 技術高中數學（第三冊）2-3 對數及其運算的意義 {status_str} ｜ PLVS 數學科  教育學習用")
    f_run.font.name = "微軟正黑體"
    f_run.font.size = Pt(9)
    f_run.font.color.rgb = RGBColor(140, 140, 140)
    
    # 4. 標題橫幅
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t_run = title_p.add_run("東大版 技術高中數學（第三冊）2-3 對數及其運算的意義 數位學習單")
    t_run.font.name = "微軟正黑體"
    t_run.font.size = Pt(16)
    t_run.bold = True
    t_run.font.color.rgb = RGBColor(30, 58, 138)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mode_text = "【教師專業教學評量與解析版 ‧ 原生 OMML 步驟算式】" if is_teacher else "【學生自主探究 ‧ 概念鷹架引導 ‧ 實戰演練學習單】"
    s_run = sub_p.add_run(mode_text)
    s_run.font.name = "微軟正黑體"
    s_run.font.size = Pt(11)
    s_run.bold = True
    s_run.font.color.rgb = RGBColor(220, 38, 38) if is_teacher else RGBColor(79, 70, 229)
    
    # 學生填寫資訊欄
    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_run = info_p.add_run("班級：___________　座號：___________　姓名：_______________　自我評量：__________ / 100")
    info_run.font.name = "微軟正黑體"
    info_run.font.size = Pt(10)
    info_run.font.color.rgb = RGBColor(71, 85, 105)
    
    # ==========================================
    # 模組一：觀念錨定與公式自我檢核
    # ==========================================
    add_heading_box(doc, "模組一：核心觀念奠基與定理填空", (22, 101, 52), "F0FDF4")
    
    add_box(
        doc,
        "重點 1：對數的定義與存在條件",
        [
            "1. 定義：設 $a > 0$、$a \\ne 1$ 且 $b > 0$，若 $a^x = b$，則以符號 $\\log_a b$ 表示 $x$，即：",
            "   $a^x = b \\iff x = \\log_a b$。稱 $a$ 為___________，稱 $b$ 為___________。",
            "2. 存在條件：底數必須滿足 $a > 0$ 且 $a \\ne 1$；真數必須滿足 $b > 0$。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "1. 稱 $a$ 為底數，稱 $b$ 為真數。",
            "2. 存在條件：底數 $a > 0, a \\ne 1$；真數 $b > 0$。"
        ] if is_teacher else None
    )
    
    add_box(
        doc,
        "重點 2：對數的五大基本運算性質（設底數與真數均滿足條件）",
        [
            "1. $\\log_a 1 = $___________，$\\log_a a = $___________。",
            "2. $\\log_{a^r} M^s = $___________$\\log_a M$（真數次方提前放分子，底數次方提前放分母）。",
            "3. $\\log_a M + \\log_a N = \\log_a ($___________$)$（同底對數相加，真數相乘）。",
            "4. $\\log_a M - \\log_a N = \\log_a ($___________$)$（同底對數相減，真數相除）。",
            "5. $a^{\\log_a M} = $___________（指數對數底數相同互相抵消）。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "1. $\\log_a 1 = 0$, $\\log_a a = 1$。",
            "2. $\\log_{a^r} M^s = \\frac{s}{r}\\log_a M$。",
            "3. $\\log_a M + \\log_a N = \\log_a (MN)$。",
            "4. $\\log_a M - \\log_a N = \\log_a (\\frac{M}{N})$。",
            "5. $a^{\\log_a M} = M$。"
        ] if is_teacher else None
    )
    
    add_box(
        doc,
        "重點 3：換底公式、連鎖律與倒數關係",
        [
            "1. 換底公式：$\\log_a b = \\frac{\\log_c b}{\\log_c a}$（其中 $c > 0, c \\ne 1$）。",
            "2. 連鎖律：$\\log_a b \\times \\log_b c \\times \\log_c d = $___________。",
            "3. 倒數關係：$\\log_a b = \\frac{1}{\\log_b a}$。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "2. 連鎖律結果為 $\\log_a d$。",
            "3. 倒數性質：$\\log_a b \\times \\log_b a = 1 \\implies \\log_a b = \\frac{1}{\\log_b a}$。"
        ] if is_teacher else None
    )

    add_box(
        doc,
        "重點 4：【名偵探破案密碼】對數神速解題口訣與記憶心法（必背錦囊）",
        [
            "1. 【旋轉蝸牛法】：$\\log_a b = x \\iff a^x = b$。從底數 $a$ 出發，逆時針繞過等號扛次方 $x$，撞向真數 $b$！",
            "2. 【四字破案訣】：真數相乘對數相加；真數相除對數相減；真上次頂底下方沉；上下同底穿透還原！",
            "3. 【高鐵換底法】：$\\log_a b = \\frac{\\log_c b}{\\log_c a}$（樓上的留在樓上，樓下的掉到樓下！底真互換直接倒數！）",
            "4. 【骨牌連鎖律】：$\\log_a b \\times \\log_b c = \\log_a c$（右上左下連連看，中間骨牌全倒光！）",
            "5. 【防錯避雷針】：內加外不加、內減外不減！即 $\\log(M+N) \\ne \\log M + \\log N$；對數相乘不能直接加！"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "★ 名偵探解題鐵律：看到對數先確認「底數大於 0 不等於 1、真數大於 0」，再靈活運用四大口訣快速化簡！"
        ] if is_teacher else None
    )
    
    # ==========================================
    # 模組二：課本核心範例與雙欄對照演練
    # ==========================================
    add_heading_box(doc, "模組二：課本精選範例與隨堂對照演練", (30, 58, 138), "EFF6FF")
    
    # 例題 1 與 隨堂 1
    add_box(
        doc,
        "【例題 1】對數的定義辨析（課本 p.75）",
        [
            "下列各式哪幾個是有意義的對數？",
            "(A) $\\log_1 5$　(B) $\\log_5 1$　(C) $\\log_{-2} 3$　(D) $\\log_2 \\frac{1}{5}$　(E) $\\log_3 (-4)$"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(A) ×：底數為 1 不合；(B) ○：底數 5、真數 1 合格；",
            "(C) ×：底數為 -2 不合；(D) ○：底數 2、真數 1/5 合格；",
            "(E) ×：真數為 -4 不合。故選 (B)、(D)。"
        ] if is_teacher else None
    )
    add_box(
        doc,
        "【隨堂練習 1】對數的定義辨析（課本 p.75）",
        [
            "下列各式哪幾個是有意義的對數？",
            "(A) $\\log_2 0$　(B) $\\log_0 2$　(C) $\\log_{\\frac{1}{2}} 5$　(D) $\\log_{\\frac{1}{2}} \\frac{5}{3}$　(E) $\\log_{-3} (-2)$"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(A) ×：真數為 0；(B) ×：底數為 0；",
            "(C) ○：底數 1/2、真數 5 合格；(D) ○：底數 1/2、真數 5/3 合格；",
            "(E) ×：底數與真數皆為負數。故選 (C)、(D)。"
        ] if is_teacher else None
    )
    
    # 例題 2 與 隨堂 2
    add_box(
        doc,
        "【例題 2】對數存在的必要條件（課本 p.76）",
        [
            "若 $\\log_{x-1}(x + 5)$ 有意義，則 $x$ 的範圍為何？"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "1. 真數條件：$x + 5 > 0 \\implies x > -5$。",
            "2. 底數條件：$x - 1 > 0$ 且 $x - 1 \\ne 1 \\implies x > 1$ 且 $x \\ne 2$。",
            "3. 取交集得：$x > 1$ 且 $x \\ne 2$。"
        ] if is_teacher else None
    )
    add_box(
        doc,
        "【隨堂練習 2】對數存在的必要條件（課本 p.76）",
        [
            "若 $\\log_{x+2}(3 - x)$ 有意義，則 $x$ 的範圍為何？"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "1. 真數條件：$3 - x > 0 \\implies x < 3$。",
            "2. 底數條件：$x + 2 > 0$ 且 $x + 2 \\ne 1 \\implies x > -2$ 且 $x \\ne -1$。",
            "3. 取交集得：$-2 < x < 3$ 且 $x \\ne -1$。"
        ] if is_teacher else None
    )
    
    # 例題 3(1~4)
    add_box(
        doc,
        "【例題 3(1)(2)】對數求值（課本 p.76）",
        [
            "(1) 試求 $\\log_2 64$ 之值。",
            "(2) 試求 $\\log_3 \\frac{1}{81}$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 因為 $2^6 = 64$，所以 $\\log_2 64 = 6$。",
            "(2) 因為 $3^{-4} = \\frac{1}{81}$，所以 $\\log_3 \\frac{1}{81} = -4$。"
        ] if is_teacher else None
    )
    add_box(
        doc,
        "【例題 3(3)(4)】對數求值（課本 p.76）",
        [
            "(3) 試求 $\\log_{0.1} 0.001$ 之值。",
            "(4) 試求 $\\log_{25} 5$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(3) 因為 $0.1^3 = 0.001$，所以 $\\log_{0.1} 0.001 = 3$。",
            "(4) 因為 $25^{\\frac{1}{2}} = \\sqrt{25} = 5$，所以 $\\log_{25} 5 = \\frac{1}{2}$。"
        ] if is_teacher else None
    )
    
    # 隨堂 3(1~4)
    add_box(
        doc,
        "【隨堂練習 3(1)(2)】對數求值（課本 p.77）",
        [
            "(1) 試求 $\\log_3 27$ 之值。",
            "(2) 試求 $\\log_2 \\frac{1}{32}$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 因為 $3^3 = 27$，故 $\\log_3 27 = 3$。",
            "(2) 因為 $2^{-5} = \\frac{1}{32}$，故 $\\log_2 \\frac{1}{32} = -5$。"
        ] if is_teacher else None
    )
    add_box(
        doc,
        "【隨堂練習 3(3)(4)】對數求值（課本 p.77）",
        [
            "(3) 試求 $\\log_{\\frac{1}{5}} \\frac{1}{625}$ 之值。",
            "(4) 試求 $\\log_{\\frac{1}{4}} 64$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(3) 因為 $(\\frac{1}{5})^4 = \\frac{1}{625}$，故值為 4。",
            "(4) 因為 $(\\frac{1}{4})^{-3} = 4^3 = 64$，故值為 -3。"
        ] if is_teacher else None
    )
    
    # 例題 4 與 隨堂 4
    add_box(
        doc,
        "【例題 4】對數基本性質求值（課本 p.79）",
        [
            "(1) 試求 $\\log_5 1 + \\log_3 3$ 之值。",
            "(2) 試求 $\\log_3 \\sqrt{3}$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\log_5 1 + \\log_3 3 = 0 + 1 = 1$。",
            "(2) $\\log_3 \\sqrt{3} = \\log_3 3^{\\frac{1}{2}} = \\frac{1}{2}\\log_3 3 = \\frac{1}{2}$。"
        ] if is_teacher else None
    )
    add_box(
        doc,
        "【隨堂練習 4】對數基本性質求值（課本 p.79）",
        [
            "(1) 試求 $\\log_{0.25} 1 + \\log_{100} 100$ 之值。",
            "(2) 試求 $\\log_{25} 125$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $0 + 1 = 1$。",
            "(2) $\\log_{5^2} 5^3 = \\frac{3}{2}\\log_5 5 = \\frac{3}{2}$。"
        ] if is_teacher else None
    )
    
    # 例題 5 與 隨堂 5
    add_box(
        doc,
        "【例題 5】對數加減運算性質（課本 p.80）",
        [
            "(1) 試求 $\\log_6 9 + \\log_6 4$ 之值。",
            "(2) 試求 $\\log_5 45 - \\log_5 9$ 之值。",
            "(3) 試求 $\\log_3 36 - \\log_3 8 + \\log_3 54$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\log_6(9 \\times 4) = \\log_6 36 = \\log_6 6^2 = 2$。",
            "(2) $\\log_5(\\frac{45}{9}) = \\log_5 5 = 1$。",
            "(3) $\\log_3(\\frac{36 \\times 54}{8}) = \\log_3(243) = \\log_3 3^5 = 5$。"
        ] if is_teacher else None
    )
    add_box(
        doc,
        "【隨堂練習 5】對數加減運算性質（課本 p.80）",
        [
            "(1) 試求 $\\log_8 16 + \\log_8 32$ 之值。",
            "(2) 試求 $\\log_3 18 - \\log_3 2$ 之值。",
            "(3) 試求 $\\log_2 12 + \\log_2 40 - \\log_2 15$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\log_8(16 \\times 32) = \\log_{2^3} 2^9 = \\frac{9}{3} = 3$。",
            "(2) $\\log_3(\\frac{18}{2}) = \\log_3 9 = 2$。",
            "(3) $\\log_2(\\frac{12 \\times 40}{15}) = \\log_2 32 = 5$。"
        ] if is_teacher else None
    )
    
    # 例題 6 與 隨堂 6
    add_box(
        doc,
        "【例題 6】換底公式與指數運算（課本 p.81）",
        [
            "(1) 試求 $\\frac{\\log_5 32}{\\log_5 2}$ 之值。",
            "(2) 試求 $5^{\\frac{\\log_2 7}{\\log_2 5}}$ 之值。",
            "(3) 試求 $9^{\\log_3 5}$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 由換底公式逆運算：$\\log_2 32 = 5$。",
            "(2) 次方 $\\frac{\\log_2 7}{\\log_2 5} = \\log_5 7$，故 $5^{\\log_5 7} = 7$。",
            "(3) $9^{\\log_3 5} = (3^2)^{\\log_3 5} = 3^{2\\log_3 5} = 3^{\\log_3 25} = 25$。"
        ] if is_teacher else None
    )
    add_box(
        doc,
        "【隨堂練習 6】換底公式與指數運算（課本 p.81）",
        [
            "(1) 試求 $\\frac{\\log_2 243}{\\log_2 9}$ 之值。",
            "(2) 試求 $4^{\\frac{\\log_7 1}{\\log_7 3}}$ 之值。",
            "(3) 試求 $3^{\\log_9 16}$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\log_9 243 = \\log_{3^2} 3^5 = \\frac{5}{2}$。",
            "(2) $\\log_7 1 = 0 \\implies 4^0 = 1$。",
            "(3) $\\log_9 16 = \\log_{3^2} 4^2 = \\log_3 4$，故 $3^{\\log_3 4} = 4$。"
        ] if is_teacher else None
    )
    
    # 例題 7 與 隨堂 7
    add_box(
        doc,
        "【例題 7】對數連鎖律（課本 p.81）",
        [
            "(1) 試求 $\\log_2 5 \\times \\log_5 2$ 之值。",
            "(2) 試求 $\\log_2 3 \\times \\log_3 5 \\times \\log_5 8$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\log_2 5 \\times \\log_5 2 = \\log_2 2 = 1$（倒數性質）。",
            "(2) 連鎖消去底真數：$\\log_2 8 = 3$。"
        ] if is_teacher else None
    )
    add_box(
        doc,
        "【隨堂練習 7】對數連鎖律（課本 p.82）",
        [
            "(1) 試求 $\\log_4 6 \\times \\log_6 4$ 之值。",
            "(2) 試求 $\\log_9 8 \\times \\log_8 7 \\times \\log_7 3$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 互為倒數相乘得 1。",
            "(2) 連鎖律得 $\\log_9 3 = \\log_{3^2} 3^1 = \\frac{1}{2}$。"
        ] if is_teacher else None
    )
    
    # 例題 8 與 隨堂 8
    add_box(
        doc,
        "【例題 8】對數性質綜合應用（課本 p.82）",
        [
            "設 $\\log_{10} 2 = x$、$\\log_{10} 3 = y$，試以 $x, y$ 表示：",
            "(1) $\\log_{10} 12$　(2) $\\log_{10} \\frac{8}{9}$　(3) $\\log_3 40$"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\log_{10}(2^2 \\times 3) = 2\\log_{10} 2 + \\log_{10} 3 = 2x + y$。",
            "(2) $\\log_{10} 8 - \\log_{10} 9 = 3x - 2y$。",
            "(3) $\\frac{\\log_{10} 40}{\\log_{10} 3} = \\frac{\\log_{10} 4 + 1}{\\log_{10} 3} = \\frac{2x + 1}{y}$。"
        ] if is_teacher else None
    )
    add_box(
        doc,
        "【隨堂練習 8】對數性質綜合應用（課本 p.82）",
        [
            "設 $\\log_{10} 3 = x$、$\\log_{10} 5 = y$，試以 $x, y$ 表示：",
            "(1) $\\log_{10} 45$　(2) $\\log_{10} \\frac{25}{27}$　(3) $\\log_{30} 15$"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\log_{10}(3^2 \\times 5) = 2x + y$。",
            "(2) $\\log_{10} 5^2 - \\log_{10} 3^3 = 2y - 3x$。",
            "(3) $\\frac{\\log_{10} 15}{\\log_{10} 30} = \\frac{\\log_{10} 3 + \\log_{10} 5}{\\log_{10} 3 + \\log_{10} 10} = \\frac{x + y}{x + 1}$。"
        ] if is_teacher else None
    )
    
    # ==========================================
    # 模組三：自我評量全套實戰挑戰
    # ==========================================
    add_heading_box(doc, "模組三：課後自我評量全套實戰挑戰（課本 p.83～p.84）", (180, 83, 9), "FFFBEB")
    
    # 觀念是非題
    add_box(
        doc,
        "【自我評量：觀念是非辨析】（請填入 ○ 或 ×）",
        [
            "1. (　) 若 $a > 0$ 且 $a \\ne 1$，則 $\\log_a 1 = 0$ 恆成立。",
            "2. (　) $\\log_a (M + N) = \\log_a M + \\log_a N$ 恆成立。",
            "3. (　) $\\frac{\\log_a M}{\\log_a N} = \\log_a (M - N)$ 恆成立。",
            "4. (　) $\\log_3 (-3)^2 = 2\\log_3 (-3)$。",
            "5. (　) 若 $b > 0$，則 $\\log_{10} b > 0$ 恆成立。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "1. ○：由 $a^0 = 1$ 定義成立。",
            "2. ×：真數相加不能拆開，公式為 $\\log_a (MN) = \\log_a M + \\log_a N$。",
            "3. ×：公式為 $\\log_a M - \\log_a N = \\log_a \\frac{M}{N}$，而 $\\frac{\\log_a M}{\\log_a N} = \\log_N M$。",
            "4. ×：左式 $\\log_3 9 = 2$，但右式真數 $-3 < 0$ 無意義。",
            "5. ×：當 $0 < b < 1$ 時，$\\log_{10} b < 0$（例如 $\\log_{10} 0.1 = -1$）。"
        ] if is_teacher else None
    )
    
    # 習題 1 & 2
    add_box(
        doc,
        "【習題 1】對數存在範圍（課本 p.83）",
        [
            "若 $\\log_{x-2}(7 - x)$ 有意義，則 $x$ 的範圍為何？"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "1. 真數條件：$7 - x > 0 \\implies x < 7$。",
            "2. 底數條件：$x - 2 > 0$ 且 $x - 2 \\ne 1 \\implies x > 2$ 且 $x \\ne 3$。",
            "3. 交集：$2 < x < 7$ 且 $x \\ne 3$。"
        ] if is_teacher else None
    )
    add_box(
        doc,
        "【習題 2】基礎對數求值（課本 p.83）",
        [
            "(1) 試求 $\\log_{10} 10 + \\log_{10} 1$ 之值。",
            "(2) 試求 $\\log_{\\sqrt{2}} 2 + \\log_9 \\sqrt{3} - \\log_{0.5} 0.125$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $1 + 0 = 1$。",
            "(2) $\\log_{2^{\\frac{1}{2}}} 2 + \\log_{3^2} 3^{\\frac{1}{2}} - \\log_{0.5} (0.5)^3 = 2 + \\frac{1}{4} - 3 = -\\frac{3}{4}$。"
        ] if is_teacher else None
    )
    
    # 習題 3 & 4
    add_box(
        doc,
        "【習題 3】對數加減運算求值（課本 p.83）",
        [
            "(1) 試求 $\\log_{12} 4 + \\log_{12} 36$ 之值。",
            "(2) 試求 $\\log_{10} 25 + \\log_{10} 12 - \\log_{10} 3$ 之值。",
            "(3) 試求 $2\\log_5 10 - 4\\log_5 2 + \\log_5 20$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\log_{12}(4 \\times 36) = \\log_{12} 144 = 2$。",
            "(2) $\\log_{10}(\\frac{25 \\times 12}{3}) = \\log_{10} 100 = 2$。",
            "(3) $\\log_5 10^2 - \\log_5 2^4 + \\log_5 20 = \\log_5(\\frac{100 \\times 20}{16}) = \\log_5 125 = 3$。"
        ] if is_teacher else None
    )
    add_box(
        doc,
        "【習題 4】指數與對數複合求值（課本 p.83）",
        [
            "(1) 試求 $4^{\\frac{\\log_7 16}{\\log_7 4}}$ 之值。",
            "(2) 試求 $3^{\\log_9 20 - \\log_9 5}$ 之值。",
            "(3) 試求 $\\log_9(\\log_2 8)$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) 次方換底：$\\log_4 16 = 2 \\implies 4^2 = 16$。",
            "(2) 次方 $\\log_9(\\frac{20}{5}) = \\log_9 4 = \\log_3 2 \\implies 3^{\\log_3 2} = 2$。",
            "(3) 括號內 $\\log_2 8 = 3 \\implies \\log_9 3 = \\log_{3^2} 3^1 = \\frac{1}{2}$。"
        ] if is_teacher else None
    )
    
    # 習題 5 & 6
    add_box(
        doc,
        "【習題 5 & 6】連鎖律與進階換底代數表示（課本 p.83～p.84）",
        [
            "(5) 試求 $\\log_9 7 \\times \\log_7 100 \\times \\log_{10} 3$ 之值。",
            "(6) 設 $a = \\log_{10} 2$、$b = \\log_{10} 3$，試以 $a, b$ 表示 $\\log_{12} 5$。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(5) $\\frac{1}{2}\\log_3 7 \\times 2\\log_7 10 \\times \\log_{10} 3 = (\\frac{1}{2} \\times 2) \\times (\\log_3 7 \\log_7 10 \\log_{10} 3) = 1 \\times 1 = 1$。",
            "(6) $\\log_{12} 5 = \\frac{\\log_{10} 5}{\\log_{10} 12} = \\frac{1 - \\log_{10} 2}{2\\log_{10} 2 + \\log_{10} 3} = \\frac{1 - a}{2a + b}$。"
        ] if is_teacher else None
    )
    
    # 習題 7
    add_box(
        doc,
        "【習題 7】常用對數近似值計算（課本 p.84）",
        [
            "已知 $\\log_{10} 2 \\approx 0.3010$、$\\log_{10} 3 \\approx 0.4771$，試求：",
            "(1) $\\log_{10} 6 \\approx $___________",
            "(2) $\\log_{10} 5 \\approx $___________",
            "(3) $\\log_{10} \\frac{3}{2} \\approx $___________"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\log_{10} 2 + \\log_{10} 3 \\approx 0.3010 + 0.4771 = 0.7781$。",
            "(2) $\\log_{10} \\frac{10}{2} = 1 - 0.3010 = 0.6990$。",
            "(3) $\\log_{10} 3 - \\log_{10} 2 \\approx 0.4771 - 0.3010 = 0.1761$。"
        ] if is_teacher else None
    )
    
    # 習題 8 & 9
    add_box(
        doc,
        "【習題 8 & 9】進階指數對數化簡（課本 p.84）",
        [
            "(8) 設 $x = \\log_{10} 5$，試求 $100^x + 10^{-x}$ 之值。",
            "(9) 試求 $(\\log_2 3 + \\log_4 9)(\\log_3 4 + \\log_9 2)$ 之值。"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(8) $100^x = (10^x)^2 = 5^2 = 25$，$10^{-x} = \\frac{1}{5} = 0.2$，總和為 $25 + 0.2 = 25.2 = \\frac{126}{5}$。",
            "(9) 前式 $\\log_2 3 + \\log_2 3 = 2\\log_2 3$；後式 $2\\log_3 2 + \\frac{1}{2}\\log_3 2 = \\frac{5}{2}\\log_3 2$；相乘得 $2 \\times \\frac{5}{2} \\times (\\log_2 3 \\log_3 2) = 5$。"
        ] if is_teacher else None
    )
    
    # 習題 10 素養生活應用
    add_box(
        doc,
        "【習題 10】生活素養應用：芮氏地震規模與能量（課本 p.84）",
        [
            "芮氏地震規模 $M$ 與釋放能量 $E$（單位：爾格）關係式為 $\\log_{10} E = 11.8 + 1.5M$：",
            "(1) 若地震波釋放 $10^{19.3}$ 爾格能量，則芮氏規模 $M$ 為何？",
            "(2) 921 集集大地震芮氏規模為 7.3，則其震央釋放之能量為多少爾格？"
        ],
        is_teacher=is_teacher,
        answer_lines=[
            "(1) $\\log_{10}(10^{19.3}) = 19.3 \\implies 19.3 = 11.8 + 1.5M \\implies 1.5M = 7.5 \\implies M = 5.0$。",
            "(2) $\\log_{10} E = 11.8 + 1.5 \\times 7.3 = 11.8 + 10.95 = 22.75 \\implies E = 10^{22.75}$ 爾格。"
        ] if is_teacher else None
    )
    
    # ==========================================
    # 模組四：學習反思與錯題自我診斷
    # ==========================================
    add_heading_box(doc, "模組四：精熟指標檢核與學習心得筆記", (79, 70, 229), "EEF2FF")
    
    add_box(
        doc,
        "【自我精熟檢核 Checklist】",
        [
            "□ 能正確寫出對數存在的底數與真數三條件（$a>0, a\\ne 1, b>0$）。",
            "□ 能熟練將對數式與指數式進行雙向互換求值。",
            "□ 能靈活運用五大基本性質，將對數的加減化為真數乘除。",
            "□ 能熟練運用換底公式與連鎖律處理異底對數問題。",
            "□ 理解常用對數在地震等現實巨量數據上的科學應用。"
        ],
        is_teacher=False
    )
    
    add_box(
        doc,
        "【課堂反思與疑惑發問區】",
        [
            "本單元我覺得最容易混淆的觀念是：_____________________________________________________",
            "我想向老師請教的一道題目是：_________________________________________________________"
        ],
        is_teacher=False
    )
    
    # 存檔
    fname = "2-3_對數及其運算的意義_教師解答版.docx" if is_teacher else "2-3_對數及其運算的意義_學生學習單.docx"
    doc.save(fname)
    print(f"[OK] 成功產出: {fname}")
    return fname

def main():
    print("[START] 開始產出 2-3 JIS B4 橫向雙欄 Word OMML 學習單")
    student_file = create_worksheet(is_teacher=False)
    teacher_file = create_worksheet(is_teacher=True)
    
    # 轉為 Base64 寫入 docx_b64.py
    with open(student_file, "rb") as f:
        s_b64 = base64.b64encode(f.read()).decode('utf-8')
    with open(teacher_file, "rb") as f:
        t_b64 = base64.b64encode(f.read()).decode('utf-8')
        
# 讀取教案 docx
    lp_file = "2-3_對數及其運算的意義_四節課教學計畫與教案.docx"
    lp_b64 = ""
    if os.path.exists(lp_file):
        with open(lp_file, "rb") as f:
            lp_b64 = base64.b64encode(f.read()).decode("utf-8")
        print(f"[OK] 教案 Base64 編碼成功: {len(lp_b64)} 字元")

    # 同步複製到 download/ 目錄
    os.makedirs("download", exist_ok=True)
    import shutil
    shutil.copy(student_file, os.path.join("download", student_file))
    shutil.copy(teacher_file, os.path.join("download", teacher_file))
    if os.path.exists(lp_file):
        shutil.copy(lp_file, os.path.join("download", lp_file))
    print("[OK] 全部 docx 已同步歸檔至 download/ 目錄")

    b64_content = f'''# -*- coding: utf-8 -*-
"""
docx_b64.py
自動生成之 Word 卷宗 Base64 資源模組 (SOP v4.0 純化旗艦標準)
"""

DOCX_STUDENT_B64 = "{s_b64}"
DOCX_TEACHER_B64 = "{t_b64}"
DOCX_LESSON_PLAN_B64 = "{lp_b64}"
'''
    with open("docx_b64.py", "w", encoding="utf-8") as f:
        f.write(b64_content)
    print("[OK] 成功產生 docx_b64.py (包含學生實作單、教師解答版、四節課教案 Base64)")

if __name__ == "__main__":
    main()
